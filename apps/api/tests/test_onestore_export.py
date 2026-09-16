"""원스토어 일괄등록 파일 추출 엔드포인트.

변환 규칙 자체는 `apps/shared/tests/utils/test_onestore_export.py` 가 지킨다.
여기서 보는 건 **엔드포인트 계약**이다:
  · 업로드 파일이 없거나 쓸 수 없으면 400 (파일을 만들지 않는다)
  · 자체 검증에 걸리면 400 + 위반 목록 (반려당하기 전에 막는다)
  · 성공하면 파일 + **무엇이 왜 빠졌는지**를 같이 돌려준다

마지막 항목이 이 엔드포인트가 바이너리가 아니라 JSON 을 돌려주는 이유다. 이 작업에서
상품이 조용히 빠지는 사고가 두 번 있었고(레거시 Play API 가 42건 누락, 국가 한 곳 때문에
36개 전부 탈락), 둘 다 화면에 안 보여서 늦게 알았다.

DB 는 in-memory SQLite. Play API 와 CDN L10N 은 네트워크라 대역으로 갈아 끼운다.
"""
import base64
import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest
from fastapi.testclient import TestClient
from shared.enums import ProductAssetUISize, ProductRarity, ProductType
from shared.models.base import Base
from shared.models.product import Category, Product
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import main
from app.api import admin as admin_api
from app.dependencies import session as session_dep
from app.utils import verify_token

URL = "/api/admin/onestore/export"

_TABLES = ("category", "category_product", "product", "price",
           "fungible_asset_product", "fungible_item_product",
           "product_gacha_entry", "product_voucher_grant")


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False},
                           poolclass=StaticPool)
    Base.metadata.create_all(
        engine, tables=[Base.metadata.tables[t] for t in _TABLES
                        if t in Base.metadata.tables])
    with Session(engine) as sess:
        yield sess


@pytest.fixture
def client(db):
    main.app.dependency_overrides[session_dep] = lambda: db
    main.app.dependency_overrides[verify_token] = lambda: None
    yield TestClient(main.app)
    main.app.dependency_overrides.clear()


def add_product(sess, sku, *, active=True, opens=None, closes=None,
                product_type=ProductType.IAP):
    category = sess.query(Category).first()
    if category is None:
        category = Category(name="Recommended", order=1, active=True,
                            l10n_key="MOBILE_SHOP_CATEGORY_Recommended")
        sess.add(category)
        sess.flush()
    product = Product(
        name=sku, order=1, google_sku=sku, product_type=product_type,
        active=active, open_timestamp=opens, close_timestamp=closes,
        rarity=ProductRarity.NORMAL, size=ProductAssetUISize.ONE_BY_ONE,
        path="=", l10n_key="=",
    )
    sess.add(product)
    sess.flush()
    category.product_list.append(product)
    sess.commit()
    return product


def onestore_export_file(rows=None):
    """개발자센터 내보내기(6열) 를 흉내낸 업로드 파일."""
    rows = rows if rows is not None else [
        ["g_pkg_old", "ko:Old|", "관리상품", "KRW : 4,400",
         "KR:KRW:4,840|US:USD:3.3|", "Android:등록중|"],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["In-App ID", "Language:In-App Title|", "In-App Type",
               "Default Price", "Country:In-App Price|", "OS:Status|"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def play_product(sku, *, kr=28000, us=(19, 990000000)):
    return {
        "productId": sku,
        "listings": [{"languageCode": "en-US", "title": f"{sku} title"}],
        "purchaseOptions": [{
            "purchaseOptionId": sku,
            "state": "ACTIVE",
            "newRegionsConfig": {"usdPrice": {"currencyCode": "USD", "units": "19",
                                              "nanos": 950000000}},
            "regionalPricingAndAvailabilityConfigs": [
                {"regionCode": "KR", "availability": "AVAILABLE",
                 "price": {"currencyCode": "KRW", "units": str(kr)}},
                {"regionCode": "US", "availability": "AVAILABLE",
                 "price": {"currencyCode": "USD", "units": str(us[0]),
                           "nanos": us[1]}},
            ],
        }],
    }


@pytest.fixture
def stub_network(monkeypatch):
    """Play API·CDN 을 대역으로. 기본은 상품 하나."""
    state = {"products": [play_product("g_pkg_a")], "l10n": {}}
    monkeypatch.setattr(admin_api, "fetch_play_onetime_products",
                        lambda *a, **k: state["products"])
    monkeypatch.setattr(admin_api, "fetch_l10n_titles", lambda *a, **k: state["l10n"])
    return state


def post(client, data=None):
    # `verify_token` 은 오버라이드하지만 라우터의 HTTPBearer 는 그대로라 헤더가 필요하다.
    return client.post(
        URL,
        files={"file": ("InAppList.xlsx",
                        data if data is not None else onestore_export_file(),
                        "application/vnd.ms-excel")},
        headers={"Authorization": "Bearer test"},
    )


def test_returns_rows_and_downloadable_file(client, db, stub_network):
    add_product(db, "g_pkg_a")

    response = post(client)

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["row_count"] == 1
    assert body["country_count"] == 2          # 업로드 파일의 KR, US
    workbook = openpyxl.load_workbook(
        io.BytesIO(base64.b64decode(body["content_base64"])))
    assert workbook.active.title == "InAppInfo"
    assert workbook.active.cell(row=2, column=1).value == "g_pkg_a"


def test_reports_already_registered_instead_of_failing(client, db, stub_network):
    """중복 ID 하나면 원스토어가 파일 전체를 반려한다. 미리 빼고, 뺐다고 알린다."""
    add_product(db, "g_pkg_old")
    add_product(db, "g_pkg_a")
    stub_network["products"] = [play_product("g_pkg_old"), play_product("g_pkg_a")]

    body = post(client).json()

    assert body["already_registered"] == ["g_pkg_old"]
    assert body["row_count"] == 1


def test_reports_skipped_with_reason(client, db, stub_network):
    """판매중인데 행이 못 된 상품은 사유와 함께 드러나야 한다."""
    add_product(db, "PLT_PACKAGE_STARTER")
    stub_network["products"] = [play_product("PLT_PACKAGE_STARTER")]

    body = post(client).json()

    assert body["row_count"] == 0
    assert body["skipped"][0]["sku"] == "PLT_PACKAGE_STARTER"
    assert "In-App ID" in body["skipped"][0]["reason"]


@pytest.mark.parametrize("opens,closes", [
    (datetime.now(timezone.utc) + timedelta(days=1), None),   # 아직 안 열림
    (None, datetime.now(timezone.utc) - timedelta(days=1)),   # 이미 닫힘
])
def test_product_outside_sale_window_is_not_exported(client, db, stub_network,
                                                     opens, closes):
    add_product(db, "g_pkg_a", opens=opens, closes=closes)

    body = post(client).json()

    assert body["row_count"] == 0
    assert body["skipped"] == []      # 범위 밖이지 문제가 아니다


def test_inactive_product_is_not_exported(client, db, stub_network):
    add_product(db, "g_pkg_a", active=False)

    assert post(client).json()["row_count"] == 0


def test_non_iap_product_is_not_exported(client, db, stub_network):
    """마일리지·무료 상품은 스토어 결제 상품이 아니다."""
    add_product(db, "g_pkg_a", product_type=ProductType.MILEAGE)

    assert post(client).json()["row_count"] == 0


def test_rejects_export_file_without_countries(client, db, stub_network):
    """등록 상품이 0개인 내보내기는 정답표가 안 된다 — 파일을 만들면 안 된다."""
    add_product(db, "g_pkg_a")

    response = post(client, onestore_export_file(rows=[]))

    assert response.status_code == 400
    assert "국가" in response.json()["detail"]


def test_rejects_non_xlsx_upload(client, db, stub_network):
    add_product(db, "g_pkg_a")

    response = post(client, b"not an excel file")

    assert response.status_code == 400


def test_rejects_when_self_validation_fails(client, db, stub_network, monkeypatch):
    """검증에 걸리면 파일을 내려주지 않는다. 반려는 원스토어가 아니라 여기서 난다."""
    add_product(db, "g_pkg_a")
    monkeypatch.setattr(admin_api, "validate_rows", lambda *a, **k: ["일부러 실패"])

    response = post(client)

    assert response.status_code == 400
    assert "일부러 실패" in response.json()["detail"]
