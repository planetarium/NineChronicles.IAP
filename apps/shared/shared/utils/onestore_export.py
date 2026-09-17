"""Google Play 인앱상품 -> 원스토어 일괄등록 파일(xlsx) 변환.

## 왜 이 모듈이 필요한가
원스토어 개발자센터의 In-App 상품 **일괄등록**에는 공식 문서가 없다. 아래 규칙은
2026-09-16 에 반려 5회로 실측해 확정한 것이고, 각 규칙의 근거가 된 반려 메시지는
테스트(`tests/utils/test_onestore_export.py`)의 docstring 에 적혀 있다.

## 진실 소스는 업로드 파일이다
배포 국가 목록도, 국가별 요구 통화도, 기본가격 통화도 문서에 없고 Play 와도 다르다
(가봉은 Play=EUR / 원스토어=USD). 개발자센터
[인앱 상품 > 관리형 상품 > 상품 일괄 등록하기 > **내보내기**] 가 주는 파일만이 안다.
그래서 이 모듈은 그 파일을 입력으로 받는다 — 상수로 박으면 배포국이 바뀔 때 조용히
틀린 파일을 만들어낸다.

## Play 쪽 함정
상품 목록은 반드시 `monetization.onetimeproducts.list` 로 받아야 한다. 레거시
`inappproducts.list` 는 상품을 빠뜨린다(151 vs 193, 빠진 42건이 전부 최근 상품).
`shared/utils/google.py` 의 `update_google_price()` 가 아직 레거시를 쓰는데, 그건
호출부가 주석처리된 죽은 코드다 — 되살릴 거면 같이 고쳐야 한다.
"""

import collections
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Mapping, Sequence

#: 템플릿. 시트명·헤더를 원스토어가 어떻게 검사하는지 모르므로 **이미 통과한 파일**을 쓴다.
TEMPLATE_PATH = Path(__file__).with_name("onestore_inapp_template.xlsx")

HEADER = ["In-App ID", "Currency:Default Price", "Country:In-App Price|",
          "Language:In-App Title|"]

#: 원스토어 In-App ID 규칙: 소문자/숫자/_/. 만, 소문자 또는 숫자로 시작, 136자 이내
ID_RULE = re.compile(r"^[a-z0-9][a-z0-9_.]{0,135}$")

#: 인앱 현지가격 허용 범위가 문서에 있는 건 이 4곳뿐이다(원스토어가 자체 스토어를
#: 운영하는 곳). 나머지는 검사하지 않는다 — 근거 없는 값으로 거르면 멀쩡한 상품이 빠진다.
#:   https://onestore-dev.gitbook.io/dev/docs/apps/product/monetization/inapp/price
PRICE_RANGE = {
    "US": (0.05, 999.99),
    "KR": (100, 600000),
    "TW": (2.00, 33000.00),
    "SG": (0.10, 1370.00),
}

#: 기본가격 통화 -> 그 통화를 쓰는 지역. Play 의 해당 지역가에서 기본가격을 만든다.
DEFAULT_CURRENCY_REGION = {"USD": "US", "KRW": "KR", "TWD": "TW", "SGD": "SG"}

#: 부가세율. 원스토어 기본가격은 **세금 미포함**인데 Play 의 한국·대만 가격은 포함가다.
VAT_RATE = {"KR": 0.10, "TW": 0.05}

#: 소수점을 쓰지 않는 통화 (ISO 4217 minor unit 0)
ZERO_DECIMAL = {"KRW", "JPY", "VND", "CLP", "PYG", "UGX", "XOF", "XAF"}

#: Play listings 의 languageCode -> 원스토어 언어코드. 순서가 출력 순서다.
LOCALE_MAP = {"ko-KR": "ko", "en-US": "en"}

#: 게임 CDN L10N CSV 컬럼 -> 원스토어 언어코드 (Play 에 제목이 없을 때의 폴백)
L10N_FALLBACK = {"ko": "Korean", "en": "English"}

#: 시즌패스는 Play 의 en-US 제목이 SKU 문자열 그대로다("COURAGEPASS34Premium").
#: 결제 화면에 노출되는 값이라 읽을 수 있는 이름으로 덮는다. 시즌 번호는 뺀다 —
#: 시즌마다 제목을 새로 넣지 않아도 되는 대신 결제 이력에서 시즌 구분은 안 된다.
TITLE_OVERRIDES = [
    ("couragepass", "Courage Pass"),
    ("adventurebosspass", "Adventure Boss Pass"),
    ("worldclearpass", "World Clear Pass"),
    ("seasonpassplus", "Season Pass Plus"),
    ("seasonpassall", "Season Pass All"),
    ("seasonpass", "Season Pass"),
]


@dataclass(frozen=True)
class OneStoreCatalog:
    """개발자센터 내보내기에서 읽은 현재 상태."""

    #: 배포 국가 -> 그 국가에 요구되는 통화. 여기 없는 국가는 넣으면 안 되고,
    #: 여기 있는 국가가 빠지면 "현지 가격이 입력되지 않았습니다" 로 반려된다.
    currency_by_country: Mapping[str, str]
    #: 이미 등록된 In-App ID. 하나라도 중복되면 **파일 전체**가 반려된다.
    registered_skus: frozenset
    #: 앱 공통 기본가격 통화(KRW/SGD/TWD/USD 중 하나).
    default_currency: str


@dataclass
class ExportResult:
    rows: list = field(default_factory=list)
    #: (sku, 사유). **비어서 빠진 상품은 여기 반드시 드러나야 한다** — 조용히 빠지는
    #: 것이 이 작업에서 가장 위험했다.
    skipped: list = field(default_factory=list)
    #: 대상 국가인데 Play 에 값이 없어 행에서 빠진 횟수 (국가별)
    uncovered_countries: dict = field(default_factory=dict)
    already_registered: list = field(default_factory=list)


def money(price: Mapping) -> float:
    """Play Money {currencyCode, units, nanos} -> float. nanos 는 없을 수 있다."""
    return int(price.get("units", 0)) + int(price.get("nanos", 0)) / 1e9


def fmt_money(amount: float, currency: str) -> str:
    """소수점 없는 통화는 정수로, 나머지는 불필요한 0 을 떼고."""
    if currency in ZERO_DECIMAL:
        return str(round(amount))
    return f"{round(amount, 2):g}"


def l10n_key(product_id: str) -> str:
    """클라이언트와 같은 규칙. IAPServiceManager.cs 의 `GetCode()` = Sku.Split("_").Last()"""
    return f"MOBILE_SHOP_PRODUCT_{product_id.split('_')[-1]}"


def _title_override(product_id: str):
    for token, name in TITLE_OVERRIDES:
        if token in product_id:
            return name
    return None


def parse_onestore_export(data: bytes) -> OneStoreCatalog:
    """내보내기 xlsx(6열) 바이트 -> 배포국·통화·기등록 SKU·기본가격 통화."""
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(c or "").strip() for c in next(rows)]
    except StopIteration:
        raise ValueError("내보내기 파일이 비어 있다")

    def column(name):
        if name not in header:
            raise ValueError(f"내보내기 파일에 '{name}' 열이 없다 (헤더: {header})")
        return header.index(name)

    id_col = column("In-App ID")
    price_col = column("Country:In-App Price|")
    default_col = column("Default Price")

    currency_by_country: dict = {}
    registered: set = set()
    default_currency = None
    for row in rows:
        if row[id_col]:
            registered.add(str(row[id_col]).strip())
        if default_currency is None and row[default_col]:
            default_currency = str(row[default_col]).split(":")[0].strip()
        for segment in str(row[price_col] or "").split("|"):
            if segment.strip():
                country, currency, _ = segment.split(":")
                currency_by_country[country.strip()] = currency.strip()

    if not currency_by_country:
        raise ValueError(
            "내보내기 파일에서 배포 국가를 찾지 못했다. 원스토어에 등록된 상품이 "
            "하나도 없으면 국가 정보가 실리지 않는다 — 상품 1개를 수기로 먼저 등록할 것"
        )
    if not default_currency:
        raise ValueError("내보내기 파일에서 기본가격 통화를 찾지 못했다")

    return OneStoreCatalog(
        currency_by_country=currency_by_country,
        registered_skus=frozenset(registered),
        default_currency=default_currency,
    )


def build_rows(play_products: Iterable[Mapping], catalog: OneStoreCatalog,
               on_sale_skus, l10n_titles: Mapping) -> ExportResult:
    """Play 상품 목록 -> 일괄등록 행.

    `on_sale_skus` 범위 밖은 **사유 없이** 건너뛴다(범위 밖이지 문제가 아니다).
    범위 안인데 행이 못 된 상품만 `skipped` 에 쌓는다.
    """
    result = ExportResult()
    uncovered: collections.Counter = collections.Counter()
    countries = sorted(catalog.currency_by_country)

    for product in sorted(play_products, key=lambda p: p.get("productId", "")):
        pid = product.get("productId", "")
        if pid not in on_sale_skus:
            continue
        if pid in catalog.registered_skus:
            result.already_registered.append(pid)
            continue

        options = product.get("purchaseOptions") or []
        if len(options) != 1:
            result.skipped.append((pid, f"purchaseOptions 가 {len(options)}개 — 수기 확인 필요"))
            continue
        option = options[0]
        if option.get("state") != "ACTIVE":
            result.skipped.append((pid, f"Play state={option.get('state')}"))
            continue
        if not ID_RULE.match(pid):
            result.skipped.append((pid, "In-App ID 규칙 위반 (소문자/숫자/_/. 만, 소문자나 숫자로 시작)"))
            continue

        regional = {
            c["regionCode"]: c
            for c in option.get("regionalPricingAndAvailabilityConfigs", [])
            if c.get("availability") == "AVAILABLE"
        }
        prices, failure, missing = _country_prices(
            regional, option, countries, catalog.currency_by_country)
        if failure:
            result.skipped.append((pid, failure))
            continue
        if not prices:
            result.skipped.append((pid, f"대상 국가 {len(countries)}곳 어디에도 가격이 없다"))
            continue

        default_price = _default_price(regional, catalog.default_currency)
        if default_price is None:
            result.skipped.append(
                (pid, f"기본가격 통화 {catalog.default_currency} 에 대응하는 지역가가 없다"))
            continue

        titles = _titles(product, l10n_titles)
        if not titles:
            result.skipped.append((pid, "Play listings·CDN L10N 양쪽에 제목이 없다"))
            continue

        # 행이 확정된 뒤에만 집계한다 — 뒤에서 제외될 상품의 빈 국가까지 세면
        # "이 국가가 비었다"가 파일에 없는 상품 얘기가 되어 오해를 부른다.
        uncovered.update(missing)
        result.rows.append([
            pid,
            f"{catalog.default_currency} : "
            f"{fmt_money(default_price, catalog.default_currency)}",
            "".join(prices),
            "".join(titles),
        ])

    result.uncovered_countries = dict(uncovered)
    return result


def _country_prices(regional, option, countries, currency_by_country):
    """국가별 현지가격 세그먼트. (세그먼트들, 치명적 사유, 값 없는 국가들)"""
    usd_fallback = None
    new_regions = option.get("newRegionsConfig") or {}
    if "usdPrice" in new_regions:
        usd_fallback = money(new_regions["usdPrice"])

    segments, missing = [], []
    for country in countries:
        want = currency_by_country[country]
        config = regional.get(country)
        if config and config["price"]["currencyCode"] == want:
            currency, amount = want, money(config["price"])
        elif want == "USD" and usd_fallback is not None:
            # Play 통화가 원스토어 요구값과 다른 23개국이 여기로 온다
            # (가봉 EUR→USD, 세네갈 XOF→USD …). EUR·XOF 는 원스토어가 받지 않는다.
            currency, amount = "USD", usd_fallback
        else:
            missing.append(country)
            continue

        low, high = PRICE_RANGE.get(country, (None, None))
        if low is not None and not (low <= amount <= high):
            return [], f"{country} 가격 {amount:g} 이 허용범위 {low}~{high} 밖", missing
        segments.append(f"{country}:{currency}:{fmt_money(amount, currency)}|")
    return segments, None, missing


def _default_price(regional, default_currency):
    region = DEFAULT_CURRENCY_REGION.get(default_currency)
    config = regional.get(region) if region else None
    if not config:
        return None
    # 원스토어 기본가격은 세금 미포함, Play 의 한국·대만 가격은 포함가다.
    return money(config["price"]) / (1 + VAT_RATE.get(region, 0.0))


def _titles(product, l10n_titles):
    play_titles = {
        listing.get("languageCode"): listing.get("title")
        for listing in product.get("listings", []) if listing.get("title")
    }
    fallback = l10n_titles.get(l10n_key(product["productId"]), {})

    override = _title_override(product["productId"])
    if override:
        resolved = {code: override for code in LOCALE_MAP.values()}
    else:
        resolved = {}
        for locale, code in LOCALE_MAP.items():
            title = play_titles.get(locale) or fallback.get(L10N_FALLBACK[code], "").strip()
            if title:
                resolved[code] = title
        # 원스토어의 '기본' 언어가 한국어라 비워두면 한국 노출 시 제목이 없는 상품이 된다.
        if "ko" not in resolved and resolved.get("en"):
            resolved["ko"] = resolved["en"]

    return [f"{code}:{resolved[code]}|"
            for code in LOCALE_MAP.values() if code in resolved]


def validate_rows(rows: Sequence[Sequence[str]], catalog: OneStoreCatalog) -> list:
    """반려당하기 전에 막는다. 위반 설명 목록을 돌려준다 (빈 목록 = 통과)."""
    violations = []

    ids = [row[0] for row in rows]
    duplicates = {i for i, n in collections.Counter(ids).items() if n > 1}
    if duplicates:
        violations.append(f"In-App ID 중복: {sorted(duplicates)}")
    already = sorted(set(ids) & set(catalog.registered_skus))
    if already:
        violations.append(f"이미 원스토어에 등록된 In-App ID: {already}")

    expected_default = f"{catalog.default_currency} : "
    for row in rows:
        pid = row[0]
        if not ID_RULE.match(pid):
            violations.append(f"{pid}: In-App ID 규칙 위반")
        if not str(row[1] or "").startswith(expected_default):
            violations.append(
                f"{pid}: 기본가격 통화가 {catalog.default_currency} 가 아니다 ({row[1]!r})")
        if not str(row[2] or "").strip():
            violations.append(f"{pid}: 국가 열이 비었다 (필수 열이다)")
            continue
        if not str(row[3] or "").strip():
            violations.append(f"{pid}: 제목이 비었다")

        got = {}
        for segment in str(row[2]).split("|"):
            if segment.strip():
                country, currency, amount = segment.split(":")
                got[country] = (currency, float(amount))

        missing = sorted(set(catalog.currency_by_country) - set(got))
        if missing:
            violations.append(f"{pid}: 배포 국가 누락 {missing}")
        extra = sorted(set(got) - set(catalog.currency_by_country))
        if extra:
            violations.append(f"{pid}: 배포 목록에 없는 국가 {extra}")
        for country, (currency, amount) in got.items():
            want = catalog.currency_by_country.get(country)
            if want and currency != want:
                violations.append(f"{pid}: {country} 통화가 {currency} — {want} 여야 한다")
            low, high = PRICE_RANGE.get(country, (None, None))
            if low is not None and not (low <= amount <= high):
                violations.append(
                    f"{pid}: {country} 가격 {amount:g} 이 허용범위 {low}~{high} 밖")

    return violations


def fetch_play_onetime_products(credential_data: str, package_name: str) -> list:
    """Play 의 관리형 상품 전량.

    ⚠️ 반드시 `monetization.onetimeproducts` 다. 레거시 `inappproducts.list` 는
    상품을 빠뜨린다(2026-09-16 실측 151 vs 193, 빠진 42건이 전부 최근 상품 —
    couragepass27~34, essone 등 **현재 판매중인 것들**). 레거시에만 있는 건 0건이라
    신규 API 가 완전한 상위집합이다.
    """
    import googleapiclient.discovery
    import json as _json

    from google.oauth2 import service_account

    credential = service_account.Credentials.from_service_account_info(
        _json.loads(credential_data),
        scopes=["https://www.googleapis.com/auth/androidpublisher"],
    )
    client = googleapiclient.discovery.build(
        "androidpublisher", "v3", credentials=credential)
    onetime = client.monetization().onetimeproducts()

    products, token = [], None
    while True:
        response = onetime.list(packageName=package_name, pageSize=100,
                                pageToken=token).execute()
        products.extend(response.get("oneTimeProducts", []))
        token = response.get("nextPageToken")
        if not token:
            return products


def fetch_l10n_titles(csv_url: str) -> dict:
    """게임 CDN 의 상품명 L10N.

    Play 의 listings 에 한국어 제목이 있는 상품은 193건 중 24건뿐이라, 원스토어의
    '기본' 언어(한국어)를 채우려면 이쪽을 같이 봐야 한다.
    """
    import csv
    import urllib.request

    with urllib.request.urlopen(csv_url, timeout=30) as response:
        body = response.read().decode("utf-8-sig")
    return {row["Key"]: row for row in csv.DictReader(io.StringIO(body))}


def write_workbook(rows: Sequence[Sequence[str]]) -> bytes:
    """템플릿에 행을 채워 xlsx 바이트로."""
    import openpyxl

    workbook = openpyxl.load_workbook(TEMPLATE_PATH)
    sheet = workbook.active
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row)  # 예시행 제거, 헤더만 남김
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
