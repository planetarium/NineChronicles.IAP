"""원스토어 인앱상품 일괄등록 파일 생성.

여기 박힌 규칙은 전부 2026-09-16 에 개발자센터 반려 5회로 **실측**해 확정한 것이다.
일괄등록에 대한 공식 문서가 없어서 다른 확인 경로가 없었다. 각 테스트의 docstring 에
어떤 반려 메시지가 그 규칙의 근거인지 적어 둔다 — 나중에 규칙이 바뀌면 그 메시지로
재현해서 확인할 수 있어야 한다.
"""

import openpyxl
import pytest

from shared.utils.onestore_export import (
    OneStoreCatalog,
    build_rows,
    parse_onestore_export,
    validate_rows,
    write_workbook,
)

# ─────────────────────────────────────────────────────────────────────────────
# 픽스처 — 실제 응답/파일에서 떠온 모양
# ─────────────────────────────────────────────────────────────────────────────


def money(currency, units, nanos=0):
    """Play 의 Money. nanos 가 없는 통화(KRW/JPY)도 있다."""
    m = {"currencyCode": currency, "units": str(units)}
    if nanos:
        m["nanos"] = nanos
    return m


def region(code, currency, units, nanos=0, availability="AVAILABLE"):
    return {
        "regionCode": code,
        "price": money(currency, units, nanos),
        "availability": availability,
    }


def play_product(product_id, *, state="ACTIVE", regions=None, listings=None,
                 usd_price=money("USD", 19, 950000000)):
    """`monetization.onetimeproducts.list` 응답 한 건."""
    option = {
        "purchaseOptionId": product_id,
        "state": state,
        "regionalPricingAndAvailabilityConfigs": regions if regions is not None else [
            region("US", "USD", 19, 990000000),
            region("KR", "KRW", 28000),
        ],
    }
    if usd_price is not None:
        option["newRegionsConfig"] = {"usdPrice": usd_price}
    return {
        "productId": product_id,
        "listings": listings if listings is not None else [
            {"languageCode": "en-US", "title": f"{product_id} EN"}
        ],
        "purchaseOptions": [option],
    }


def catalog(currency_by_country=None, registered=(), default_currency="KRW"):
    return OneStoreCatalog(
        currency_by_country=currency_by_country or {"US": "USD", "KR": "KRW"},
        registered_skus=frozenset(registered),
        default_currency=default_currency,
    )


def export_workbook(rows):
    """개발자센터 [상품 일괄 등록하기 > 내보내기] 가 주는 6열 xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["In-App ID", "Language:In-App Title|", "In-App Type",
               "Default Price", "Country:In-App Price|", "OS:Status|"])
    for r in rows:
        ws.append(r)
    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# 업로드 파일 파싱 — 배포국·통화·기등록 SKU 의 유일한 진실 소스
# ─────────────────────────────────────────────────────────────────────────────


def test_parse_export_yields_country_currency_map():
    """배포 국가와 국가별 통화는 문서에 없고 Play 와도 다르다.

    가봉(GA)은 Play 가 EUR 로 값을 매기는데 원스토어는 USD 를 요구한다. 그래서
    Play 통화를 그대로 쓰면 "지원하지 않는 국가코드-통화코드 조합입니다" 로 반려됐다.
    내보내기 파일만이 정답을 안다.
    """
    data = export_workbook([
        ["g_single_ap01", "ko:AP Potion 15|", "관리상품", "KRW : 4,400",
         "GA:USD:3.77|KR:KRW:4,840|JP:JPY:660|", "Android:등록중|"],
    ])

    result = parse_onestore_export(data)

    assert result.currency_by_country == {"GA": "USD", "KR": "KRW", "JP": "JPY"}


def test_parse_export_yields_registered_skus():
    """중복 In-App ID 가 하나라도 있으면 파일 전체가 반려된다.

    반려 메시지: "입력하신 In-App ID 는 이미 등록되어 있습니다."
    같은 파일이 기등록 목록도 들고 있으므로 업로드 한 번으로 중복 제외까지 해결된다.
    """
    data = export_workbook([
        ["g_single_ap01", "ko:AP|", "관리상품", "KRW : 4,400", "US:USD:3.3|", "Android:등록중|"],
        ["g_pkg_wc1", "ko:WC|", "관리상품", "KRW : 40,000", "US:USD:30|", "Android:등록중|"],
    ])

    result = parse_onestore_export(data)

    assert result.registered_skus == frozenset({"g_single_ap01", "g_pkg_wc1"})


def test_parse_export_yields_default_currency():
    """기본가격 통화는 앱 전체 공통 설정이라 추측하면 안 된다.

    반려 메시지: "ONEconsole에서 설정된 기본가격 통화와 다른 통화코드가 입력 되었습니다."
    내보내기의 Default Price 열이 그 값을 그대로 알려준다.
    """
    data = export_workbook([
        ["g_single_ap01", "ko:AP|", "관리상품", "KRW : 4,400", "US:USD:3.3|", "Android:등록중|"],
    ])

    assert parse_onestore_export(data).default_currency == "KRW"


def test_parse_export_rejects_file_without_countries():
    """등록된 상품이 0개면 국가 정보가 없어 정답표를 만들 수 없다.

    닭-달걀이라 막다른 길로 보이지만, 상품 하나를 수기 등록하면 풀린다.
    조용히 빈 매핑을 돌려주면 그 뒤 모든 행이 국가 없이 만들어져 반려된다.
    """
    data = export_workbook([])

    with pytest.raises(ValueError, match="국가"):
        parse_onestore_export(data)


# ─────────────────────────────────────────────────────────────────────────────
# 행 생성
# ─────────────────────────────────────────────────────────────────────────────


def test_row_uses_play_price_when_currency_matches():
    """통화가 맞는 국가는 Play 가격을 그대로 써서 가격 포인트($19.99)를 보존한다.

    국가는 알파벳순으로 나열한다 — 실제로 통과한 파일이 그 순서였다.
    """
    result = build_rows([play_product("g_pkg_a")], catalog(), {"g_pkg_a"}, {})

    assert result.rows[0][2] == "KR:KRW:28000|US:USD:19.99|"


def test_row_falls_back_to_usd_when_play_currency_differs():
    """Play 통화가 원스토어 요구값과 다르면 newRegionsConfig.usdPrice 로 채운다.

    23개국이 여기 해당한다(가봉 EUR→USD, 세네갈 XOF→USD, UAE AED→USD …).
    EUR·XOF·AED 는 원스토어가 아예 받지 않는다.
    """
    product = play_product("g_pkg_a", regions=[
        region("US", "USD", 19, 990000000),
        region("GA", "EUR", 17, 160000000),
    ])

    result = build_rows([product], catalog({"US": "USD", "GA": "USD"},
                                           default_currency="USD"),
                        {"g_pkg_a"}, {})

    assert result.rows[0][2] == "GA:USD:19.95|US:USD:19.99|"


def test_product_without_default_currency_region_is_reported():
    """기본가격은 앱 공통 통화(예: KRW)의 지역가에서 만든다. 그 지역가가 없으면
    기본가격을 만들 수 없으므로 사유를 남기고 뺀다."""
    product = play_product("g_pkg_a", regions=[region("US", "USD", 19, 990000000)])

    result = build_rows([product], catalog({"US": "USD"}), {"g_pkg_a"}, {})

    assert result.rows == []
    assert "기본가격" in result.skipped[0][1]


def test_default_price_excludes_vat():
    """기본가격은 세금 미포함인데 Play 의 한국 가격은 부가세 포함가다.

    검산: 개발자센터에 등록된 g_pkg_worldclearpass1premium 이 기본가격 40,000 /
    KR 현지가 44,000 이고, Play 의 KR 가격이 44,000 이다. 즉 기본가격 = Play KR ÷ 1.1.
    """
    product = play_product("g_pkg_a", regions=[
        region("US", "USD", 29, 990000000),
        region("KR", "KRW", 44000),
    ])

    result = build_rows([product], catalog(), {"g_pkg_a"}, {})

    assert result.rows[0][1] == "KRW : 40000"


def test_already_registered_sku_is_excluded():
    data = catalog(registered={"g_pkg_a"})

    result = build_rows([play_product("g_pkg_a"), play_product("g_pkg_b")],
                        data, {"g_pkg_a", "g_pkg_b"}, {})

    assert [r[0] for r in result.rows] == ["g_pkg_b"]
    assert result.already_registered == ["g_pkg_a"]


def test_product_not_on_sale_is_excluded_silently():
    """판매중이 아닌 상품은 제외 사유로 쌓지 않는다 — 범위 밖이지 문제가 아니다."""
    result = build_rows([play_product("g_pkg_a"), play_product("g_pkg_b")],
                        catalog(), {"g_pkg_a"}, {})

    assert [r[0] for r in result.rows] == ["g_pkg_a"]
    assert result.skipped == []


@pytest.mark.parametrize("product,reason", [
    (play_product("g_pkg_a", state="INACTIVE_PUBLISHED"), "state"),
    (play_product("PLT_PACKAGE_STARTER"), "In-App ID"),
    (play_product("g_pkg_a", listings=[]), "제목"),
])
def test_unusable_product_is_reported_not_dropped(product, reason):
    """빠진 상품은 반드시 사유와 함께 드러나야 한다.

    조용히 빠지는 게 이 작업에서 제일 위험했다 — 레거시 Play API 가 42건을 말없이
    누락했고, 앙길라 한 곳 때문에 36개가 통째로 사라진 적도 있다.
    """
    result = build_rows([product], catalog(), {product["productId"]}, {})

    assert result.rows == []
    assert len(result.skipped) == 1
    assert reason in result.skipped[0][1]


def test_missing_country_price_is_counted_not_fatal():
    """대상 국가 중 일부만 없으면 그 국가만 빠지고 상품은 남는다."""
    product = play_product("g_pkg_a", regions=[region("US", "USD", 19, 990000000)],
                           usd_price=None)

    result = build_rows([product], catalog({"US": "USD", "AI": "XCD"},
                                           default_currency="USD"),
                        {"g_pkg_a"}, {})

    assert result.rows[0][2] == "US:USD:19.99|"
    assert result.uncovered_countries == {"AI": 1}


def test_season_pass_title_is_cleaned():
    """Play 의 시즌패스 제목이 SKU 문자열 그대로라 결제 화면에 그게 노출된다."""
    product = play_product("g_pkg_couragepass34premium", listings=[
        {"languageCode": "en-US", "title": "COURAGEPASS34Premium"}])

    result = build_rows([product], catalog(), {product["productId"]}, {})

    assert result.rows[0][3] == "ko:Courage Pass|en:Courage Pass|"


def test_korean_title_falls_back_to_cdn_then_english():
    """원스토어의 '기본' 언어가 한국어라 비워두면 한국 노출 시 제목이 없다.

    Play 에 ko-KR 제목이 있는 상품은 193건 중 24건뿐이라 게임 CDN L10N 을 먼저 보고,
    그것도 없으면 영어로 채운다(운영 결정).
    """
    l10n = {"MOBILE_SHOP_PRODUCT_a": {"Korean": "가나다", "English": "ABC"}}

    with_cdn = build_rows([play_product("g_pkg_a")], catalog(), {"g_pkg_a"}, l10n)
    without_cdn = build_rows([play_product("g_pkg_b")], catalog(), {"g_pkg_b"}, {})

    assert with_cdn.rows[0][3] == "ko:가나다|en:g_pkg_a EN|"
    assert without_cdn.rows[0][3] == "ko:g_pkg_b EN|en:g_pkg_b EN|"


def test_play_korean_title_wins_over_cdn():
    product = play_product("g_pkg_a", listings=[
        {"languageCode": "en-US", "title": "ABC"},
        {"languageCode": "ko-KR", "title": "플레이 제목"},
    ])
    l10n = {"MOBILE_SHOP_PRODUCT_a": {"Korean": "CDN 제목", "English": "ABC"}}

    result = build_rows([product], catalog(), {"g_pkg_a"}, l10n)

    assert result.rows[0][3] == "ko:플레이 제목|en:ABC|"


# ─────────────────────────────────────────────────────────────────────────────
# 자체 검증 — 반려당하기 전에 막는다
# ─────────────────────────────────────────────────────────────────────────────


def test_validate_passes_on_good_rows():
    data = catalog()
    result = build_rows([play_product("g_pkg_a")], data, {"g_pkg_a"}, {})

    assert validate_rows(result.rows, data) == []


def test_validate_rejects_row_missing_a_country():
    """배포 국가가 하나라도 빠지면 "<국가> : 현지 가격이 입력되지 않았습니다" 로 반려된다."""
    data = catalog({"US": "USD", "KR": "KRW", "JP": "JPY"})
    rows = [["g_pkg_a", "KRW : 100", "US:USD:1|KR:KRW:1100|", "ko:A|en:A|"]]

    violations = validate_rows(rows, data)

    assert any("JP" in v for v in violations)


def test_validate_rejects_wrong_currency():
    data = catalog({"US": "USD", "JP": "JPY"})
    rows = [["g_pkg_a", "KRW : 100", "US:USD:1|JP:USD:1|", "ko:A|en:A|"]]

    violations = validate_rows(rows, data)

    assert any("JP" in v and "USD" in v for v in violations)


def test_validate_rejects_default_currency_mismatch():
    data = catalog(default_currency="KRW")
    rows = [["g_pkg_a", "USD : 1", "US:USD:1|KR:KRW:1100|", "ko:A|en:A|"]]

    assert any("기본가격" in v for v in validate_rows(rows, data))


def test_validate_rejects_duplicate_id():
    data = catalog()
    row = ["g_pkg_a", "KRW : 100", "US:USD:1|KR:KRW:1100|", "ko:A|en:A|"]

    assert any("중복" in v for v in validate_rows([row, list(row)], data))


def test_validate_rejects_empty_country_column():
    """국가 열을 비우면 "잘못된 값 입니다" 로 전 행이 반려된다. 필수 열이다."""
    data = catalog()
    rows = [["g_pkg_a", "KRW : 100", "", "ko:A|en:A|"]]

    assert any("국가" in v for v in validate_rows(rows, data))


def test_validate_rejects_price_out_of_documented_range():
    """문서화된 4개국(US/KR/TW/SG)만 범위를 안다. 나머지는 검사하지 않는다."""
    data = catalog({"US": "USD", "KR": "KRW"})
    rows = [["g_pkg_a", "KRW : 100", "US:USD:0.01|KR:KRW:1100|", "ko:A|en:A|"]]

    assert any("US" in v for v in validate_rows(rows, data))


# ─────────────────────────────────────────────────────────────────────────────
# 파일 쓰기
# ─────────────────────────────────────────────────────────────────────────────


def test_workbook_keeps_template_shape():
    """시트명과 헤더를 원스토어가 어떻게 검사하는지 모른다.

    이미 통과한 템플릿을 그대로 쓰면 그 불확실성이 사라진다.
    """
    rows = [["g_pkg_a", "KRW : 100", "US:USD:1|", "ko:A|en:A|"]]

    wb = openpyxl.load_workbook(__import__("io").BytesIO(write_workbook(rows)))
    ws = wb.active

    assert ws.title == "InAppInfo"
    assert list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0] == (
        "In-App ID", "Currency:Default Price", "Country:In-App Price|",
        "Language:In-App Title|")
    assert list(ws.iter_rows(min_row=2, values_only=True)) == [tuple(rows[0])]
